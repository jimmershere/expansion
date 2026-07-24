from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from mscout.models import Listing

# Requested fields first, decision-support fields after.
COLUMNS = [
    ("Make", 14),
    ("Model", 20),
    ("Year", 7),
    ("Mileage", 10),
    ("Clean Carfax", 12),
    ("URL", 46),
    ("Avg Resale Price (Northern US)", 16),
    ("Asking Price", 12),
    ("Below Market $", 13),
    ("Below Market %", 13),
    ("Location", 18),
    ("Category", 16),
    ("Comps Used", 11),
    ("Notes", 50),
]

_HEADER_FILL = PatternFill("solid", fgColor="1F3B57")
_DEAL_FILL = PatternFill("solid", fgColor="E7F4E4")


def write_xlsx(listings: list[Listing], path: str, deal_threshold_pct: float) -> None:
    listings = sorted(
        listings, key=lambda l: (l.below_market_pct is None, -(l.below_market_pct or 0))
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Deals"
    ws.freeze_panes = "A2"

    for col, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width

    for row, l in enumerate(listings, start=2):
        notes = "; ".join([l.title_notes] + l.notes) if (l.title_notes or l.notes) else ""
        values = [
            l.make or "?", l.model or l.title[:28], l.year, l.mileage,
            l.clean_carfax, l.url, l.avg_resale_north, l.asking_price,
            l.below_market, l.below_market_pct, l.location, l.category,
            l.comp_count or None, notes,
        ]
        for col, v in enumerate(values, start=1):
            ws.cell(row=row, column=col, value=v)
        for col_name, fmt in (("Avg Resale Price (Northern US)", '"$"#,##0'),
                              ("Asking Price", '"$"#,##0'),
                              ("Below Market $", '"$"#,##0'),
                              ("Below Market %", '0.0"%"'),
                              ("Mileage", "#,##0")):
            idx = [c[0] for c in COLUMNS].index(col_name) + 1
            ws.cell(row=row, column=idx).number_format = fmt
        if (l.below_market_pct or 0) >= deal_threshold_pct:
            for col in range(1, len(COLUMNS) + 1):
                ws.cell(row=row, column=col).fill = _DEAL_FILL

    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(len(listings) + 1, 2)}"
    wb.save(path)
